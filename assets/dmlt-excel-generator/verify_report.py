"""
Recalculate a generated DMLT workbook and assert it behaves dynamically.
=======================================================================

openpyxl writes formulas but never evaluates them, so a workbook can look
correct and still be full of ``#REF!`` / ``#NAME?`` once opened.  This script
closes that gap by driving a real spreadsheet engine (LibreOffice, over UNO)
to do a full recalculation, and then checks two things:

  1. **Zero formula errors** — no cell evaluates to ``#REF!``, ``#VALUE!``,
     ``#NAME?``, ``#DIV/0!``, ``#N/A``, ``#NULL!`` or ``#NUM!``.
  2. **The workbook is actually dynamic** — flipping ``Include?`` from ``Y``
     to ``N`` on MAIN must move the HARDWARE_SIZING and SUMMARY totals.  If
     the numbers were Python-computed literals they would not budge.

Usage
-----
    python verify_report.py report.xlsx
    python verify_report.py report.xlsx --flip 500   # rows to switch to "N"

Requires LibreOffice (``soffice``) and the ``uno`` Python bindings.
Exit code is 0 only when both checks pass.
"""

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
import time

import openpyxl

#: Values Excel/Calc render for a failed formula.
ERROR_VALUES = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")

_SOFFICE_PORT = 2002


# ---------------------------------------------------------------------------
# LibreOffice / UNO plumbing
# ---------------------------------------------------------------------------

def _path_to_url(path: str) -> str:
    import uno
    return uno.systemPathToFileUrl(os.path.abspath(path))


def _start_soffice(profile_dir: str):
    """Launch a headless soffice listening on a UNO socket. Returns the Popen."""
    cmd = [
        "soffice",
        "--headless", "--norestore", "--nologo", "--nodefault",
        f"-env:UserInstallation={_path_to_url(profile_dir)}",
        f"--accept=socket,host=127.0.0.1,port={_SOFFICE_PORT};urp;",
    ]
    return subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def _connect(timeout: int = 120):
    """Connect to the running soffice instance, retrying until it is up."""
    import uno
    ctx_local = uno.getComponentContext()
    resolver = ctx_local.ServiceManager.createInstanceWithContext(
        "com.sun.star.bridge.UnoUrlResolver", ctx_local
    )
    url = (
        f"uno:socket,host=127.0.0.1,port={_SOFFICE_PORT};"
        "urp;StarOffice.ComponentContext"
    )
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        try:
            return resolver.resolve(url)
        except Exception as exc:                     # noqa: BLE001 — retry loop
            last = exc
            time.sleep(1)
    raise RuntimeError(f"Could not connect to soffice: {last}")


def recalculate(in_path: str, out_path: str) -> None:
    """
    Open ``in_path``, force a full recalculation, save the result to
    ``out_path`` as .xlsx with cached values.
    """
    import uno

    profile = tempfile.mkdtemp(prefix="dmlt_lo_profile_")
    proc = _start_soffice(profile)
    try:
        ctx = _connect()
        desktop = ctx.ServiceManager.createInstanceWithContext(
            "com.sun.star.frame.Desktop", ctx
        )

        def prop(name, value):
            p = uno.createUnoStruct("com.sun.star.beans.PropertyValue")
            p.Name, p.Value = name, value
            return p

        doc = desktop.loadComponentFromURL(
            _path_to_url(in_path), "_blank", 0,
            (prop("Hidden", True), prop("UpdateDocMode", 3)),
        )
        try:
            doc.calculateAll()
            doc.storeToURL(
                _path_to_url(out_path),
                (prop("FilterName", "Calc MS Excel 2007 XML"),),
            )
        finally:
            doc.close(False)
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


# ---------------------------------------------------------------------------
# Inspection
# ---------------------------------------------------------------------------

def scan_errors(path: str) -> list:
    """Return ``[(sheet, coord, value), ...]`` for every error cell."""
    wb = openpyxl.load_workbook(path, data_only=True)
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() in ERROR_VALUES:
                    found.append((ws.title, cell.coordinate, v.strip()))
    wb.close()
    return found


def read_totals(path: str) -> dict:
    """Pull the headline numbers the dynamic check compares before/after."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}

    hw = wb["HARDWARE_SIZING"]
    for row in hw.iter_rows(min_col=2, max_col=5):
        if str(row[0].value).strip().upper() == "TOTAL":
            out["hw_total_source_mb"] = row[1].value
            out["hw_total_target_mb"] = row[2].value
            out["hw_total_target_gb"] = row[3].value
            break

    sm = wb["SUMMARY"]
    wanted = {
        "Included Source Size (GB)": "summary_included_gb",
        "Target S/4HANA Size (GB)": "summary_target_gb",
        "Tables Included": "summary_tables_included",
        "Tables Excluded": "summary_tables_excluded",
    }
    for row in sm.iter_rows(min_col=2, max_col=3):
        label = str(row[0].value).strip() if row[0].value else ""
        if label in wanted:
            out[wanted[label]] = row[1].value

    wb.close()
    return out


def flip_include(path: str, out_path: str, count: int) -> int:
    """
    Set the first ``count`` MAIN ``Include?`` cells that currently hold "Y"
    to "N".  Returns how many were flipped.
    """
    wb = openpyxl.load_workbook(path)
    ws = wb["MAIN"]

    header_row = None
    include_col = None
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Include?":
                header_row, include_col = cell.row, cell.column
                break
        if include_col:
            break
    if not include_col:
        raise RuntimeError("Could not locate the Include? column on MAIN")

    flipped = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if flipped >= count:
            break
        cell = ws.cell(row=r, column=include_col)
        if isinstance(cell.value, str) and cell.value.strip().upper() == "Y":
            cell.value = "N"
            flipped += 1

    wb.save(out_path)
    wb.close()
    return flipped


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook", help="Generated .xlsx to verify")
    ap.add_argument("--flip", type=int, default=500,
                    help="How many Include? cells to switch to N (default 500)")
    ap.add_argument("--keep", metavar="DIR",
                    help="Keep the recalculated intermediates in DIR")
    args = ap.parse_args()

    work = args.keep or tempfile.mkdtemp(prefix="dmlt_verify_")
    os.makedirs(work, exist_ok=True)

    baseline = os.path.join(work, "recalc_baseline.xlsx")
    flipped_src = os.path.join(work, "flipped_src.xlsx")
    flipped = os.path.join(work, "recalc_flipped.xlsx")

    print(f"[1/4] Recalculating {args.workbook} …", flush=True)
    recalculate(args.workbook, baseline)

    print("[2/4] Scanning for formula errors …", flush=True)
    errors = scan_errors(baseline)
    if errors:
        print(f"  FAIL — {len(errors)} formula error cell(s):")
        for sheet, coord, val in errors[:25]:
            print(f"    {sheet}!{coord} = {val}")
        if len(errors) > 25:
            print(f"    … and {len(errors) - 25} more")
    else:
        print("  PASS — zero formula errors")

    before = read_totals(baseline)
    print("  Baseline totals:")
    for k, v in before.items():
        print(f"    {k:28} = {v}")

    print(f"[3/4] Flipping {args.flip} Include? cells to N …", flush=True)
    n = flip_include(args.workbook, flipped_src, args.flip)
    print(f"  flipped {n} cell(s)")

    print("[4/4] Recalculating the flipped workbook …", flush=True)
    recalculate(flipped_src, flipped)
    after = read_totals(flipped)
    print("  Totals after flip:")
    for k, v in after.items():
        print(f"    {k:28} = {v}")

    flip_errors = scan_errors(flipped)

    # ---- verdict ----------------------------------------------------------
    def moved(key):
        a, b = before.get(key), after.get(key)
        return isinstance(a, (int, float)) and isinstance(b, (int, float)) and a != b

    dynamic_keys = [
        "hw_total_source_mb", "hw_total_target_mb",
        "summary_included_gb", "summary_tables_included",
        "summary_tables_excluded",
    ]
    changed = [k for k in dynamic_keys if moved(k)]

    print("\n=== VERDICT ===")
    ok_errors = not errors and not flip_errors
    print(f"  formula errors (baseline)  : {len(errors)}")
    print(f"  formula errors (flipped)   : {len(flip_errors)}")
    print(f"  totals that responded      : {changed or 'NONE'}")

    ok_dynamic = len(changed) >= 3
    if ok_errors and ok_dynamic:
        print("  RESULT: PASS — workbook recalculates cleanly and is fully dynamic")
        return 0

    if not ok_errors:
        print("  RESULT: FAIL — formula errors present")
    if not ok_dynamic:
        print("  RESULT: FAIL — totals did not respond to Include? changes "
              "(values are probably hardcoded)")
    return 1


if __name__ == "__main__":
    sys.exit(main())
