"""
GROWTH sheet — doc counts per table × year, summed across systems.
Rows = tables, Columns = years ascending, last column = Total.
Year 0 is excluded.
"""
from collections import defaultdict
from openpyxl.utils import get_column_letter
from styles import (
    setup_sheet, write_title_banner, write_section_header,
    style_cell, fill_formula, fill_colhdr, fill_section,
    font_body, font_colhdr, font_formula, font_section,
    align_left, align_center, align_right, border_thin,
    FMT_INT, C_WHITE
)
from openpyxl.styles import Font, Alignment


def build(ws, growth_data: list):
    setup_sheet(ws)

    # --- Aggregate: {tabname: {year: total_doc_count}}
    agg: dict[str, dict[int, int]] = defaultdict(lambda: defaultdict(int))
    for r in growth_data:
        yr = r.get("gjahr")
        if not yr or yr == 0:
            continue
        agg[r["tabname"]][int(yr)] += int(r.get("doc_count") or 0)

    tables = sorted(agg.keys())
    years  = sorted({yr for t in agg.values() for yr in t.keys()})

    # Column layout: A=spacer, B=table, C..N=years, last=Total
    YEAR_START_COL = 3  # C
    total_col      = YEAR_START_COL + len(years)   # after all year cols

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 24
    for i in range(len(years) + 1):
        ws.column_dimensions[get_column_letter(YEAR_START_COL + i)].width = 14

    end_col = total_col

    next_row = write_title_banner(
        ws, row=1, col_start=2, col_end=end_col,
        title="DOCUMENT GROWTH",
        subtitle=(
            "Annual document counts per table summed across all source systems. "
            "Use to identify steep-growth tables and archiving candidates."
        ),
        title_size=16, row_height=30
    )
    next_row += 1
    next_row = write_section_header(ws, next_row, 2, end_col,
                                    "YEAR-ON-YEAR DOCUMENT COUNTS (all systems combined)")

    # --- Column headers
    ws.cell(row=next_row, column=2, value="Table").fill      = fill_colhdr
    ws.cell(row=next_row, column=2).font      = Font(name="Arial", bold=True,
                                                     color=C_WHITE, size=9)
    ws.cell(row=next_row, column=2).alignment = align_left
    ws.cell(row=next_row, column=2).border    = border_thin

    for i, yr in enumerate(years):
        col = YEAR_START_COL + i
        c = ws.cell(row=next_row, column=col, value=str(yr))
        style_cell(c, fill=fill_colhdr,
                   font=Font(name="Arial", bold=True, color=C_WHITE, size=9),
                   alignment=align_center, border=border_thin)

    c_total = ws.cell(row=next_row, column=total_col, value="Total")
    style_cell(c_total, fill=fill_colhdr,
               font=Font(name="Arial", bold=True, color=C_WHITE, size=9),
               alignment=align_center, border=border_thin)

    ws.row_dimensions[next_row].height = 16
    ws.freeze_panes = f"B{next_row + 1}"
    next_row += 1

    # --- Data rows
    for table in tables:
        year_vals = agg[table]
        row_values = [year_vals.get(yr, 0) for yr in years]
        row_total  = sum(row_values)

        c = ws.cell(row=next_row, column=2, value=table)
        style_cell(c, font=font_body(9, bold=True),
                   alignment=align_left, border=border_thin)

        for i, val in enumerate(row_values):
            col = YEAR_START_COL + i
            c = ws.cell(row=next_row, column=col, value=val if val else None)
            style_cell(c, font=font_body(9),
                       alignment=align_right, border=border_thin,
                       number_format=FMT_INT)

        c_t = ws.cell(row=next_row, column=total_col, value=row_total)
        style_cell(c_t, fill=fill_formula, font=font_body(9, bold=True),
                   alignment=align_right, border=border_thin,
                   number_format=FMT_INT)

        ws.row_dimensions[next_row].height = 15
        next_row += 1
