"""
COLLISIONS sheet — company-code collisions detected deterministically.

Detection rule:
  From org rows where from_type="BUKRS" AND to_type="CLIENT"
  group by from_id.
  A collision = same from_id in >1 sidclnt with a DIFFERENT from_text.
"""
from dmlt_calc import detect_collisions   # shared calculation engine
from openpyxl.utils import get_column_letter
from styles import (
    setup_sheet, write_title_banner, write_section_header, write_col_headers,
    style_cell, fill_col, fill_formula, fill_section,
    font_body, font_col, font_formula, font_section,
    align_left, align_center, align_right, border_thin,
    C_COL_FONT, C_WHITE, FMT_INT
)
from openpyxl.styles import Font, PatternFill, Alignment


def build(ws, org_data: list) -> int:
    """Build the COLLISIONS sheet.  Returns collision count."""
    setup_sheet(ws)

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16   # Company Code
    ws.column_dimensions["C"].width = 18   # System-Client
    ws.column_dimensions["D"].width = 40   # Entity Name
    ws.column_dimensions["E"].width = 30   # Risk
    ws.column_dimensions["F"].width = 14

    collisions = detect_collisions(org_data)

    subtitle = (
        f"{len(collisions)} collision(s) detected — "
        "same company-code number, different entity names across systems. "
        "Each must be renumbered before consolidation."
    )
    next_row = write_title_banner(
        ws, row=1, col_start=2, col_end=6,
        title="COMPANY-CODE COLLISIONS",
        subtitle=subtitle,
        title_size=16, row_height=30
    )
    next_row += 1

    if not collisions:
        ws.merge_cells(f"B{next_row}:F{next_row}")
        c = ws.cell(row=next_row, column=2,
                    value="No company-code collisions detected in this run.")
        c.font      = font_body(10, bold=True, color="2E7D32")
        c.alignment = align_left
        return 0

    # Summary KPI row
    ws.merge_cells(f"B{next_row}:C{next_row}")
    kpi_c = ws.cell(row=next_row, column=2,
                    value=f"Total Collisions: {len(collisions)}")
    kpi_c.fill      = fill_col
    kpi_c.font      = Font(name="Arial", bold=True, size=11, color=C_COL_FONT)
    kpi_c.alignment = align_center
    kpi_c.border    = border_thin
    ws.row_dimensions[next_row].height = 20
    next_row += 2

    for coll in collisions:
        # Block header — company code
        next_row = write_section_header(
            ws, next_row, 2, 6,
            f"Company Code: {coll['bukrs']}   |   "
            f"{len(coll['systems'])} systems  |  Risk: HIGH — must renumber on consolidation"
        )

        # Column headers for this block
        next_row = write_col_headers(
            ws, next_row, 2,
            ["Company Code", "System-Client", "Entity Name", "Risk"],
            [16, 18, 40, 30]
        )

        for sys_row in coll["systems"]:
            ws.cell(row=next_row, column=2, value=coll["bukrs"])
            ws.cell(row=next_row, column=3, value=sys_row["sidclnt"])
            ws.cell(row=next_row, column=4, value=sys_row["name"])
            ws.cell(row=next_row, column=5,
                    value="HIGH — renumber on merge")

            for col in range(2, 6):
                c = ws.cell(row=next_row, column=col)
                style_cell(c,
                           fill=fill_col,
                           font=font_col(9),
                           alignment=align_left,
                           border=border_thin)
            ws.row_dimensions[next_row].height = 15
            next_row += 1

        next_row += 1   # blank row between blocks

    return len(collisions)
