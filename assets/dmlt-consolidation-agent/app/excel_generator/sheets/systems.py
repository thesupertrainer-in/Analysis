"""
SYSTEMS sheet — one row per sidclnt from the system section.
"""
from openpyxl.styles import Alignment
from styles import (
    setup_sheet, write_title_banner, write_section_header, write_col_headers,
    style_cell, fill_formula, fill_input,
    font_body, font_formula,
    align_left, align_center, align_right, border_thin,
    FMT_INT
)


HEADERS = [
    "System-Client", "SID", "DB System", "Release",
    "SP Level", "Unicode", "Occupied Vol (GB)", "Total Vol (GB)"
]
WIDTHS = [14, 10, 12, 10, 10, 10, 18, 16]


def build(ws, system_data: list):
    setup_sheet(ws)

    # Column layout: A=spacer, B..I = data
    ws.column_dimensions["A"].width = 3
    for i, w in enumerate(WIDTHS):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i + 2)].width = w

    next_row = write_title_banner(
        ws, row=1, col_start=2, col_end=9,
        title="SYSTEMS",
        subtitle="Source system profiles extracted from the DMLT extractor",
        title_size=16, row_height=30
    )
    next_row += 1
    next_row = write_section_header(ws, next_row, 2, 9, "SOURCE SYSTEM INVENTORY")
    next_row = write_col_headers(ws, next_row, 2, HEADERS, WIDTHS)

    ws.freeze_panes = f"B{next_row}"

    for row_data in system_data:
        occ = row_data.get("occupied_vol") or 0
        tot = row_data.get("total_vol") or 0
        unicode_val = "Yes" if row_data.get("unicode") == "X" else "No"

        values = [
            row_data.get("sidclnt", ""),
            row_data.get("sid", ""),
            row_data.get("db_sys", ""),
            row_data.get("rel", ""),
            row_data.get("sp_level", ""),
            unicode_val,
            round(occ / 1024, 2) if occ else occ,
            round(tot / 1024, 2) if tot else tot,
        ]

        aligns = [
            align_left, align_left, align_center, align_center,
            align_center, align_center, align_right, align_right
        ]
        fills  = [None, None, None, None, None, None, fill_formula, fill_formula]
        fmts   = [None, None, None, None, None, None, FMT_INT, FMT_INT]

        for i, (val, aln, fil, fmt) in enumerate(
                zip(values, aligns, fills, fmts)):
            c = ws.cell(row=next_row, column=i + 2, value=val)
            style_cell(c,
                       fill=fil,
                       font=font_formula(9) if fil == fill_formula else font_body(9),
                       alignment=aln,
                       border=border_thin,
                       number_format=fmt)
        ws.row_dimensions[next_row].height = 15
        next_row += 1
