"""
MAIN_DECISION — the control grid the agent seeds from the consultant's answers.

Systems are rows, the categories actually present in the run are columns, and
every cell is an editable "Y"/"N".  The agent fills this in at generation time
and writes the same decisions into MAIN's Include? cells.

There is deliberately no live link from this grid to MAIN: no macros, no
formulas pointing back.  MAIN is the source of truth once the workbook is
open — this sheet records what the run was generated with, so a reader can
see the starting position at a glance.  Editing it after the fact does not
move MAIN, which is why the subtitle says so on the sheet itself.
"""
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font

from styles import (
    setup_sheet, style_cell,
    fill_title, fill_subbanner, fill_colhdr, fill_input, fill_rowlbl,
    font_title, font_subbanner, font_colhdr, font_data,
    align_left, align_center, border_thin,
    C_INC_FONT, C_WHITE,
)

TITLE    = "CONSOLIDATION DECISION GRID"
SUBTITLE = ("Agent fills from your answers · Y = migrate, N = leave behind · "
            "MAIN is the source of truth")

ROW_TITLE = 1
ROW_SUB   = 2
ROW_HDR   = 4
ROW_DATA  = 5

WIDTH_SYSTEM = 18
WIDTH_CAT    = 11


def build(ws, sidcltns: list, categories: list,
          decisions: dict = None, base_system: str = None) -> dict:
    """
    sidcltns    – ordered system-client ids, one row each
    categories  – category codes present in the data, one column each
    decisions   – {sidclnt: {category: "Y"|"N"}}; anything missing defaults "Y"
    base_system – sidclnt chosen as the migration shell; labelled "(base)".
                  This is a migration decision made by the consultant, not
                  something derived from volume, so it arrives as a parameter.

    Returns a refs dict.
    """
    setup_sheet(ws)
    decisions = decisions or {}

    n_cols = 1 + len(categories)
    last_letter = get_column_letter(n_cols)

    ws.column_dimensions["A"].width = WIDTH_SYSTEM
    for i in range(len(categories)):
        ws.column_dimensions[get_column_letter(2 + i)].width = WIDTH_CAT

    # ── banner ──────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{ROW_TITLE}:{last_letter}{ROW_TITLE}")
    style_cell(ws.cell(row=ROW_TITLE, column=1, value=TITLE),
               fill=fill_title, font=font_title(16),
               alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[ROW_TITLE].height = 30

    ws.merge_cells(f"A{ROW_SUB}:{last_letter}{ROW_SUB}")
    style_cell(ws.cell(row=ROW_SUB, column=1, value=SUBTITLE),
               fill=fill_subbanner, font=font_subbanner(9),
               alignment=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[ROW_SUB].height = 15

    # ── header row ──────────────────────────────────────────────────────────
    style_cell(ws.cell(row=ROW_HDR, column=1, value="System"),
               fill=fill_colhdr, font=font_colhdr(9),
               alignment=align_center, border=border_thin)
    for i, cat in enumerate(categories):
        style_cell(ws.cell(row=ROW_HDR, column=2 + i, value=cat),
                   fill=fill_colhdr, font=font_colhdr(9),
                   alignment=align_center, border=border_thin)
    ws.row_dimensions[ROW_HDR].height = 15

    # ── one row per system ──────────────────────────────────────────────────
    fnt_cell = Font(name="Arial", size=10, bold=True, color=C_INC_FONT)
    row = ROW_DATA
    for sid in sidcltns:
        label = f"{sid} (base)" if base_system and sid == base_system else str(sid)
        style_cell(ws.cell(row=row, column=1, value=label),
                   fill=fill_rowlbl, font=font_data(9, bold=True),
                   alignment=align_left, border=border_thin)
        for i, cat in enumerate(categories):
            val = (decisions.get(sid, {}) or {}).get(cat, "Y")
            style_cell(ws.cell(row=row, column=2 + i, value=val),
                       fill=fill_input, font=fnt_cell,
                       alignment=align_center, border=border_thin)
        ws.row_dimensions[row].height = 15
        row += 1

    return {
        "data_start_row": ROW_DATA,
        "data_end_row":   row - 1,
        "categories":     list(categories),
        "sidcltns":       list(sidcltns),
        "base_system":    base_system,
    }
