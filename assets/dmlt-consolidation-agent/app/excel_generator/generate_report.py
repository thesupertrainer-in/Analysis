"""
DMLT Consolidation Study — Excel Report Generator
==================================================

Usage (standalone):
    python generate_report.py \\
        --master  path/to/master.json \\
        --growth  path/to/growth.json \\
        --system  path/to/system.json \\
        --org     path/to/org.json \\
        --nriv    path/to/nriv.json \\
        --out     output/report.xlsx

The script is engagement-agnostic: it discovers sidclnt values, categories,
years, and system counts from the data — nothing is hardcoded.

Principles:
  - Deterministic detection and pivoting in Python (data rows only).
  - Every calculated cell is an Excel formula — never a Python-computed
    literal written to a cell.
  - The workbook recalculates correctly when a user edits Include? in MAIN.

Sheets produced: COVER -> SYSTEMS -> MAIN_DECISION -> MAIN.
"""

import argparse
import json
import os
import sys

# ── make local sheets/ package importable regardless of CWD ─────────────────
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import openpyxl

from sheets.cover         import build as build_cover
from sheets.systems       import build as build_systems
from sheets.main_decision import build as build_main_decision
from sheets.main_sheet    import build as build_main

# NOTE — sheets beyond MAIN are intentionally not built in this pass.
# sheets/collisions.py, conflicts.py, growth.py, hardware_sizing.py,
# summary.py and dashboard.py are still on disk, but they read MAIN through
# the old flat layout (one system column, one Include? column, one Target MB
# column).  The pivoted MAIN has none of those, so wiring them up now would
# produce a sizing model that is quietly wrong.  The sizing model is being
# redesigned against the pivoted MAIN in the next pass; these builders get
# reconnected there.


# ---------------------------------------------------------------------------
# JSON loading helpers
# ---------------------------------------------------------------------------

def _load_section(path: str, expected_section: str) -> list:
    """Load a single-section JSON file and return its data array."""
    with open(path, "r", encoding="utf-8") as fh:
        doc = json.load(fh)

    section = doc.get("section", "")
    if section and section.lower() != expected_section.lower():
        print(
            f"[WARN] File {path!r}: section={section!r}, "
            f"expected {expected_section!r} — proceeding anyway."
        )

    data = doc.get("data", [])
    if not isinstance(data, list):
        raise ValueError(
            f"File {path!r}: 'data' field must be a JSON array, "
            f"got {type(data).__name__}."
        )
    return data


def load_sections(paths: dict) -> dict:
    """
    paths = { "master": "...", "growth": "...", "system": "...",
              "org": "...", "nriv": "..." }
    Returns the same dict with list values instead of path strings.
    """
    return {sec: _load_section(p, sec) for sec, p in paths.items()}


# ---------------------------------------------------------------------------
# Main generator
# ---------------------------------------------------------------------------

def generate_report(sections: dict, out_path: str,
                    base_system: str = None,
                    decisions: dict = None) -> str:
    """
    Build the COVER / SYSTEMS / MAIN_DECISION / MAIN workbook.

    Parameters
    ----------
    sections : dict
        Keys: master, growth, system, org, nriv — each a list of row dicts.
        Only master and system are read in this pass; growth, org and nriv are
        accepted so the call signature stays stable for the next pass.
    out_path : str
        Destination .xlsx file path (parent directory is created if needed).
    base_system : str, optional
        The sidclnt chosen as the migration shell.  Labelled "(base)" on
        MAIN_DECISION.  This is a consultant decision, not something derived
        from volume, so the caller supplies it; the agent asks the user and
        passes it through.  Unknown or omitted values simply mean no row is
        marked.
    decisions : dict, optional
        {sidclnt: {category: "Y"|"N"}} seed for the decision grid.  Anything
        not named defaults to "Y".  Written into MAIN's Include? cells too, so
        the two sheets agree at generation time.

    Returns
    -------
    str — absolute path of the written .xlsx file.
    """
    master_data = sections["master"]
    system_data = sections["system"]

    # Systems come from the system section when present, else from master —
    # either way the count is discovered, never assumed.
    sidcltns = sorted({r["sidclnt"] for r in system_data}) or sorted(
        {r.get("sidclnt", "") for r in master_data if r.get("sidclnt")}
    )
    categories = sorted({
        str(r.get("category") or "").strip()
        for r in master_data
        if str(r.get("agg_level", "")).upper() == "T" and r.get("category")
    })

    if base_system and base_system not in sidcltns:
        print(f"[WARN] base system {base_system!r} is not one of {sidcltns} — "
              f"no row will be marked (base).")

    run_meta = {
        "sidcltns":           sidcltns,
        "num_source_systems": len(sidcltns),
    }

    # ── Create workbook ────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── Sheet 1: COVER ─────────────────────────────────────────────────────
    ws_cover = wb.create_sheet("COVER")
    build_cover(ws_cover, run_meta)

    # ── Sheet 2: SYSTEMS ───────────────────────────────────────────────────
    ws_systems = wb.create_sheet("SYSTEMS")
    build_systems(ws_systems, system_data)

    # ── Sheet 3: MAIN_DECISION ─────────────────────────────────────────────
    ws_dec = wb.create_sheet("MAIN_DECISION")
    dec_refs = build_main_decision(
        ws_dec,
        sidcltns=sidcltns,
        categories=categories,
        decisions=decisions,
        base_system=base_system,
    )

    # ── Sheet 4: MAIN ──────────────────────────────────────────────────────
    ws_main = wb.create_sheet("MAIN")
    main_refs = build_main(
        ws_main,
        master_data=master_data,
        sidcltns=sidcltns,
        decisions=decisions,
    )

    # ── Force recalculation on open ────────────────────────────────────────
    # openpyxl writes formulas but never their cached results, so every
    # calculated cell is empty until something evaluates it.  Setting this flag
    # makes Excel/Calc do a full recalculation the moment the file opens,
    # instead of showing blanks to whoever opens it first.
    wb.calculation.fullCalcOnLoad = True

    # ── Write file ─────────────────────────────────────────────────────────
    out_path = os.path.abspath(out_path)
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    print(f"[OK] Workbook written: {out_path}")
    print(f"     Sheets    : {[w.title for w in wb.worksheets]}")
    print(f"     Systems   : {sidcltns}")
    print(f"     Categories: {categories}")
    print(f"     Base      : {base_system or '(none set)'}")
    print(f"     Tables    : {main_refs['table_count']:,} "
          f"(MAIN rows {main_refs['data_start_row']}:{main_refs['data_end_row']})")
    print(f"     Grid      : {len(dec_refs['sidcltns'])} systems "
          f"x {len(dec_refs['categories'])} categories")

    return out_path


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def _cli():
    parser = argparse.ArgumentParser(
        description="Generate the DMLT Consolidation Study Excel report."
    )
    parser.add_argument("--master",  required=True, help="Path to master JSON file")
    parser.add_argument("--growth",  required=True, help="Path to growth JSON file")
    parser.add_argument("--system",  required=True, help="Path to system JSON file")
    parser.add_argument("--org",     required=True, help="Path to org JSON file")
    parser.add_argument("--nriv",    required=True, help="Path to nriv JSON file")
    parser.add_argument("--out",     default="output/dmlt_report.xlsx",
                        help="Output .xlsx path (default: output/dmlt_report.xlsx)")
    parser.add_argument("--base-system", default=None, metavar="SIDCLNT",
                        help="System chosen as the migration shell; labelled "
                             "'(base)' on MAIN_DECISION (e.g. RQ1_500)")
    args = parser.parse_args()

    paths = {
        "master": args.master,
        "growth": args.growth,
        "system": args.system,
        "org":    args.org,
        "nriv":   args.nriv,
    }

    print("[INFO] Loading JSON sections …")
    sections = load_sections(paths)
    for sec, data in sections.items():
        print(f"       {sec}: {len(data):,} rows")

    generate_report(sections, args.out, base_system=args.base_system)


if __name__ == "__main__":
    _cli()
